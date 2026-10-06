#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_demo_templates.py — RayLaw v2 架空デモ書式 20 種（T01..T20）の生成スクリプト

生成物（templates/demo/<kind>/）
  template.docx       python-docx で生成。A4 縦・単一セクション・横書き・Noto Serif JP 12pt。
                      ヘッダー {document_status_label}、フッター PAGE フィールド + {template_version_label}。
                      本文は docxtemplater の allowlist タグのみ（{key} / {#key}..{/key} / {^key}..{/key}）。
  schema.json         TemplateFieldSchemaSchema（packages/template-schema/src/fieldSpec.ts）準拠。
  provenance.json     出典・適用範囲・確認日。sampleOnly:true、supervised:false（監修済みとは表示しない）。
  sample_values.json  架空のデモ値（FieldValues 形 {key:{rawValue}} + attachments + renderContext）。
  ast_sample.json     本文 DocumentAST（schemaVersion 1, blocks: heading/numberedParagraph/paragraph）。
  ../README.md        20 書式の一覧。

実行:  python3 -I scripts/build_demo_templates.py [--out DIR] [--check]
  冪等: 生成バイト列が既存ファイルと一致する場合は書き換えない（ZIP のタイムスタンプを固定し、
  entry 順を揃えて再圧縮するため、同じ入力からは同じ bytes が得られる）。
  --check: 書き込まずに差分の有無を報告する（差分があれば exit 1）。

依存: python-docx 1.2.x, lxml。TypeScript パッケージには触れない。
"""
from __future__ import annotations

import argparse
import datetime as _dt
import hashlib
import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path
from typing import Any, Callable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "templates" / "demo"

FONT = "Noto Serif JP"
CHECKED_AT = "2026-10-06"
ISSUE_DATE_ISO = "2026-10-06"
ISSUE_DATE_WAREKI = "令和８年１０月６日"
FIXED_DT = _dt.datetime(2026, 10, 6, 0, 0, 0)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS

TAG_KEY_RE = re.compile(r"^[a-z][a-z0-9_]{0,63}$")
DANGEROUS_KEYS = {"__proto__", "constructor", "prototype", "this", "eval"}

# packages/template-schema/src/validate.ts の SYSTEM_SUPPLIED_KEYS と同じ集合（スキーマで定義しなくてよいキー）
SYSTEM_SUPPLIED_KEYS = {
    "document_status_label", "template_version_label", "matter_id", "issue_date", "issue_date_wareki",
    "body_text", "body_blocks", "text", "number_label", "is_heading", "is_numbered", "is_paragraph",
    "has_attachments", "attachments", "name", "copies",
    "office_name", "office_address", "office_postcode", "office_phone", "office_fax", "office_email",
    "lawyer_name", "lawyer_title", "matter_title",
    "recipient_name", "recipient_honorific", "recipient_organization", "recipient_address",
    "client_name", "opponent_name", "subject",
}

# ---------------------------------------------------------------------------
# 架空デモデータ（実在の人物・団体とは無関係）
# ---------------------------------------------------------------------------
OFFICE = {
    "name": "青葉法律事務所",
    "postcode": "１００－００１３",
    "address": "東京都千代田区霞が関一丁目１番１号　青葉デモビル５階",
    "phone": "０３－００００－０００１",
    "fax": "０３－００００－０００２",
    "email": "demo@aoba-law.example",
}
LAWYER = {"name": "青葉　花子", "title": "弁護士"}
MATTER = {
    "title": "売買代金請求事件（株式会社ミドリ商事）",
    "case_number": "令和８年（ワ）第１２３４号",
    "court": "東京地方裁判所",
    "court_department": "民事第１２部",
    "scope": "髙﨑産業株式会社に対する売買代金請求に関する交渉及び訴訟（サンプル）",
    "client_name": "株式会社ミドリ商事",
    "opponent_name": "髙﨑産業株式会社",
}
CLIENT_REP = "代表取締役　緑川　太郎"
OPPONENT_REP = "代表取締役　髙﨑　誠"
CLIENT_ADDRESS = "東京都中央区日本橋デモ町一丁目２番３号"
OPPONENT_ADDRESS = "埼玉県さいたま市大宮区デモ本町四丁目５番６号　髙﨑ビル２階"

# 同姓同名の相手方代理人 2 名（所属が異なる）
AGENT_A = {"name": "佐藤　一郎", "organization": "さくら総合法律事務所", "address": "東京都港区虎ノ門二丁目２番２号　虎ノ門デモビル３階", "postcode": "１０５－０００１", "honorific": "先生"}
AGENT_B = {"name": "佐藤　一郎", "organization": "港中央法律事務所", "address": "神奈川県横浜市中区日本大通１２番地　港デモ会館７階", "postcode": "２３１－００２１", "honorific": "先生"}
COURT_RCPT = {"name": "東京地方裁判所民事第１２部", "organization": "東京地方裁判所", "address": "東京都千代田区霞が関一丁目１番４号", "postcode": "１００－８９２０", "honorific": "御中"}
CLIENT_RCPT = {"name": "株式会社ミドリ商事", "organization": "株式会社ミドリ商事　代表取締役　緑川　太郎", "address": CLIENT_ADDRESS, "postcode": "１０３－００２７", "honorific": "御中"}
OPPONENT_RCPT = {"name": "髙﨑産業株式会社", "organization": "髙﨑産業株式会社　代表取締役　髙﨑　誠", "address": OPPONENT_ADDRESS, "postcode": "３３０－０８４６", "honorific": "御中"}
# 30 grapheme 以上の長い名称（段落伸長の検証用）
LONG_ORG_NAME = "一般社団法人全国中小企業取引適正化推進協議会関東甲信越支部事務局"
THIRD_RCPT = {"name": "渡邊　﨑太郎", "organization": LONG_ORG_NAME, "address": "東京都新宿区西新宿デモ二丁目８番１号　デモタワー３０階", "postcode": "１６３－８００１", "honorific": "様"}

SAMPLE_NOTE = "架空のデモ値。実在の人物・団体・事件とは無関係。監修済みではない（sample_only）。"
DEMO_DISCLAIMER = "【サンプル文】本書は RayLaw のデモ用に作成した架空の文例です。実際の事実関係・主張・法的評価は担当弁護士が確認のうえ記載してください。"


def date_value(raw: str, iso: str, precision: str = "day") -> dict[str, str]:
    return {"raw": raw, "iso": iso, "precision": precision}


def money_value(raw: str, decimal: str) -> dict[str, str]:
    return {"raw": raw, "decimal": decimal}


# ---------------------------------------------------------------------------
# FieldSpec ヘルパー
# ---------------------------------------------------------------------------
def src(key: str, source: str, path: str, label: str, type_: str = "text", required: bool = False, **kw: Any) -> dict[str, Any]:
    f: dict[str, Any] = {"key": key, "type": type_, "label": label, "source": source, "sourcePath": path, "required": required, "defaultPolicy": "from_source"}
    if type_ == "multiline":
        f["outputFormat"] = "multiline"
    f.update(kw)
    return f


def inp(key: str, label: str, type_: str = "text", required: bool = False, **kw: Any) -> dict[str, Any]:
    f: dict[str, Any] = {"key": key, "type": type_, "label": label, "source": "input", "required": required, "defaultPolicy": "none"}
    if type_ == "multiline":
        f["outputFormat"] = "multiline"
    if type_ == "date":
        f["outputFormat"] = "plain"
        f["validationRules"] = [{"rule": "isoDateOrRaw"}]
        f.setdefault("description", "原文（raw）を保持し、そのまま差し込む。ISO 値は解析候補として別に保持する。")
    if type_ == "money":
        f["outputFormat"] = "plain"
        f["validationRules"] = [{"rule": "jpyInteger"}]
        f.setdefault("description", "金額原文（raw）をそのまま差し込む。解析値（JPY 整数 decimal string）は別に保持し、値を黙って変更しない。")
    f.update(kw)
    return f


def legal(key: str, label: str, type_: str, **kw: Any) -> dict[str, Any]:
    """法的判断を含む項目: legalJudgement:true, defaultPolicy none（既定値を置かない）。"""
    f = inp(key, label, type_, **kw)
    f["legalJudgement"] = True
    f["defaultPolicy"] = "none"
    return f


def arr(key: str, label: str, items: list[dict[str, Any]], required: bool = False, max_items: int = 100, **kw: Any) -> dict[str, Any]:
    f: dict[str, Any] = {"key": key, "type": "array", "label": label, "source": "input", "required": required, "defaultPolicy": "none", "itemFields": items, "maxItems": max_items}
    f.update(kw)
    return f


def item(key: str, label: str, type_: str = "text", required: bool = False, options: list[str] | None = None, max_graphemes: int | None = None) -> dict[str, Any]:
    it: dict[str, Any] = {"key": key, "type": type_, "label": label, "required": required}
    if options is not None:
        it["options"] = [{"value": o, "label": o} for o in options]
    if max_graphemes:
        it["maxGraphemes"] = max_graphemes
    return it


def opts(*values: str) -> list[dict[str, str]]:
    return [{"value": v, "label": v} for v in values]


OFFICE_FIELDS = [
    src("office_name", "office", "office.name", "自事務所名", required=True),
    src("office_postcode", "office", "office.postcode", "事務所郵便番号"),
    src("office_address", "office", "office.address", "事務所所在地"),
    src("office_phone", "office", "office.phone", "事務所電話番号"),
    src("office_fax", "office", "office.fax", "事務所ＦＡＸ番号"),
    src("lawyer_name", "office", "lawyer.name", "担当弁護士氏名", required=True),
]
RECIPIENT_FIELDS = [
    src("recipient_organization", "recipient", "recipient.organization", "宛先の所属・組織"),
    src("recipient_name", "recipient", "recipient.name", "宛名", required=True, validationRules=[{"rule": "noDoubleHonorific"}], description="敬称を含めない。敬称は recipient_honorific（所内設定）で付加し、二重敬称を検出する。"),
    src("recipient_honorific", "recipient", "recipient.honorific", "敬称", description="人物「先生/様」、組織「御中」等を所内設定から選択。"),
]
MATTER_CORE = [
    src("matter_title", "matter", "matter.title", "案件名", required=True),
    src("client_name", "matter", "matter.client_name", "依頼者", required=True),
    src("opponent_name", "matter", "matter.opponent_name", "相手方"),
]
SUBJECT = inp("subject", "件名", required=True, maxGraphemes=100)
CASE_NUMBER = src("case_number", "matter", "matter.case_number", "事件番号（原文）", required=True, description="先頭 0 や全角数字を含む文字列として保持する。")
COURT_NAME = src("court_name", "matter", "matter.court", "裁判所", required=True)
COURT_DEPT = src("court_department", "matter", "matter.court_department", "係属部")
PLAINTIFF = inp("plaintiff_name", "原告（表示）", required=True, description="依頼者/相手方のどちらが原告かは案件ごとに人が指定する。氏名一致だけで置換しない。")
DEFENDANT = inp("defendant_name", "被告（表示）", required=True)
SUBMITTER_POSITION = inp("submitter_position", "提出者の立場", "enum", required=True, options=opts("原告", "被告", "申立人", "相手方", "その他"))
EVIDENCE_REFS = arr("evidence_refs", "証拠参照（証拠方法）", [item("label", "証拠番号（表示ラベル）", required=True, max_graphemes=50), item("title", "標目")], max_items=200)

# サンプル値（共通）
OFFICE_VALUES = {"office_name": OFFICE["name"], "office_postcode": OFFICE["postcode"], "office_address": OFFICE["address"], "office_phone": OFFICE["phone"], "office_fax": OFFICE["fax"], "lawyer_name": LAWYER["name"]}
MATTER_VALUES = {"matter_title": MATTER["title"], "client_name": MATTER["client_name"], "opponent_name": MATTER["opponent_name"]}


def recipient_values(r: dict[str, str]) -> dict[str, str]:
    return {"recipient_organization": r["organization"], "recipient_name": r["name"], "recipient_honorific": r["honorific"]}


def court_caption_values() -> dict[str, Any]:
    return {"case_number": MATTER["case_number"], "court_name": MATTER["court"], "court_department": MATTER["court_department"]}


def rv(values: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """plain dict → FieldValues 形 {key: {rawValue}}"""
    return {k: {"rawValue": v} for k, v in values.items()}


# ---------------------------------------------------------------------------
# AST ヘルパー
# ---------------------------------------------------------------------------
class AstBuilder:
    def __init__(self) -> None:
        self.blocks: list[dict[str, Any]] = []

    def _id(self) -> str:
        return f"b{len(self.blocks) + 1}"

    def heading(self, text: str, level: int = 2) -> "AstBuilder":
        self.blocks.append({"id": self._id(), "type": "heading", "level": level, "text": text})
        return self

    def numbered(self, label: str, text: str, level: int = 1) -> "AstBuilder":
        self.blocks.append({"id": self._id(), "type": "numberedParagraph", "numberLabel": label, "text": text, "level": level})
        return self

    def para(self, text: str) -> "AstBuilder":
        self.blocks.append({"id": self._id(), "type": "paragraph", "text": text})
        return self

    def build(self) -> dict[str, Any]:
        return {"schemaVersion": 1, "blocks": self.blocks}


def ast_lines(ast: dict[str, Any]) -> list[dict[str, Any]]:
    """renderData.ts の body_blocks と同じ形へ（検証用レンダリングで使用）"""
    out = []
    for b in ast["blocks"]:
        t = b["type"]
        out.append({
            "text": b.get("text", b.get("displayLabel", "")),
            "number_label": b.get("numberLabel", ""),
            "is_heading": t == "heading",
            "is_numbered": t == "numberedParagraph",
            "is_paragraph": t in ("paragraph", "evidenceRef"),
        })
    return out


# ---------------------------------------------------------------------------
# python-docx ヘルパー
# ---------------------------------------------------------------------------
def _set_rfonts(rpr: Any, font: str = FONT) -> None:
    """w:rFonts の ascii/hAnsi/eastAsia/cs を明示し、テーマ参照属性を取り除く。"""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), font)
    rfonts.set(qn("w:hint"), "eastAsia")


def _set_lang(rpr: Any) -> None:
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "ja-JP")
    lang.set(qn("w:eastAsia"), "ja-JP")
    lang.attrib.pop(qn("w:bidi"), None)


def _strip_optional_parts(doc: Any) -> None:
    """python-docx 既定テンプレートの thumbnail / customXml / stylesWithEffects を外し、最小構成にする。"""
    pkg = doc.part.package
    for rid, rel in list(pkg.rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            del pkg.rels[rid]
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/customXml") or rel.reltype.endswith("/stylesWithEffects"):
            del doc.part.rels[rid]


def _configure_styles(doc: Any) -> None:
    styles_el = doc.styles.element
    rpr_default = styles_el.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _set_rfonts(rpr_default)
    for tag in ("w:sz", "w:szCs"):
        el = rpr_default.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr_default.append(el)
        el.set(qn("w:val"), "24")
    _set_lang(rpr_default)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    _set_rfonts(normal.element.get_or_add_rPr())
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.line_spacing = Pt(20)

    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    _set_rfonts(h2.element.get_or_add_rPr())
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(2)
    h2.paragraph_format.keep_with_next = True

    for name in ("Header", "Footer"):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(9)
        _set_rfonts(st.element.get_or_add_rPr())


def _set_core_props(doc: Any, title: str, kind: str) -> None:
    cp = doc.core_properties
    cp.author = "RayLaw demo builder"
    cp.last_modified_by = "RayLaw demo builder"
    cp.title = f"{title}（架空デモ書式 {kind}）"
    cp.subject = "RayLaw demo template sample_only"
    cp.description = "架空のデモ書式。監修済みではない。所内弁護士の確認と承認前は正式出力不可。"
    cp.keywords = "raylaw;demo;sample_only"
    cp.category = "demo"
    cp.comments = ""
    cp.created = FIXED_DT
    cp.modified = FIXED_DT
    cp.revision = 1


def _add_page_field(paragraph: Any, size_pt: int = 9) -> None:
    """静的ページ番号フィールド（allowlist）: <w:fldSimple w:instr="PAGE">"""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size_pt * 2))
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def _write_header_footer(doc: Any) -> None:
    sec = doc.sections[0]
    hdr = sec.header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("{document_status_label}")
    hr.bold = True
    hr.font.size = Pt(10)

    ftr = sec.footer
    ftr.is_linked_to_previous = False
    fp1 = ftr.paragraphs[0]
    fp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp1.add_run("－ ")
    r.font.size = Pt(9)
    _add_page_field(fp1)
    r = fp1.add_run(" －")
    r.font.size = Pt(9)
    fp2 = ftr.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp2.add_run("{template_version_label}")
    fr.font.size = Pt(8)


def new_document(title: str, kind: str) -> Any:
    doc = Document()
    _strip_optional_parts(doc)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(30)
    sec.bottom_margin = Mm(25)
    sec.left_margin = Mm(25)
    sec.right_margin = Mm(25)
    sec.header_distance = Mm(12)
    sec.footer_distance = Mm(12)
    _configure_styles(doc)
    _set_core_props(doc, title, kind)
    _write_header_footer(doc)
    # 既定テンプレートの先頭空段落を取り除く
    body = doc.element.body
    for p in list(body.findall(qn("w:p"))):
        body.remove(p)
    return doc


def para(doc: Any, text: str = "", *, align: Any = None, bold: bool = False, size: int | None = None, style: str | None = None,
         indent_left: Any = None, first_line: Any = None, space_before: Any = None, space_after: Any = None) -> Any:
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if text:
        r = p.add_run(text)
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    if indent_left is not None:
        pf.left_indent = indent_left
    if first_line is not None:
        pf.first_line_indent = first_line
    if space_before is not None:
        pf.space_before = space_before
    if space_after is not None:
        pf.space_after = space_after
    return p


RIGHT = WD_ALIGN_PARAGRAPH.RIGHT
CENTER = WD_ALIGN_PARAGRAPH.CENTER


def split_tag_paragraph(doc: Any, prefix: str, tag: str, suffix: str = "", *, split_at: int | None = None, align: Any = None) -> Any:
    """タグを意図的に複数 w:r に分割する（書式違いの run）。例: '{recip'(bold) + 'ient_name}'(normal)。"""
    assert tag.startswith("{") and tag.endswith("}")
    if split_at is None:
        split_at = max(2, len(tag) // 2)
    p = doc.add_paragraph()
    if prefix:
        p.add_run(prefix)
    r1 = p.add_run(tag[:split_at])
    r1.bold = True
    r2 = p.add_run(tag[split_at:])
    r2.bold = False
    r2.italic = False
    if suffix:
        p.add_run(suffix)
    if align is not None:
        p.alignment = align
    return p


def loop_block(doc: Any, key: str, inner: Callable[[], None], inverted: bool = False) -> None:
    """段落ループ: 開始タグだけの段落 / 内容段落 / 終了タグだけの段落。docxtemplater(paragraphLoop) が段落単位で繰り返す。"""
    para(doc, ("{^%s}" if inverted else "{#%s}") % key)
    inner()
    para(doc, "{/%s}" % key)


def add_body_blocks(doc: Any) -> None:
    """本文 AST → 見出し/番号付き/通常段落。各条件もタグだけの段落で囲み、未使用の空段落を残さない。"""
    para(doc, "{#body_blocks}")
    para(doc, "{#is_heading}")
    hp = para(doc, style="Heading 2")
    hr = hp.add_run("{text}")
    hr.bold = True
    para(doc, "{/is_heading}")
    para(doc, "{#is_numbered}")
    para(doc, "{number_label}　{text}", indent_left=Pt(12))
    para(doc, "{/is_numbered}")
    para(doc, "{#is_paragraph}")
    para(doc, "{text}", first_line=Pt(12))
    para(doc, "{/is_paragraph}")
    para(doc, "{/body_blocks}")


def add_attachments_block(doc: Any, heading: str = "添付書類") -> None:
    para(doc, "{#has_attachments}")
    para(doc, heading, bold=True, space_before=Pt(8))
    para(doc, "{#attachments}")
    para(doc, "{name}　{copies}", indent_left=Pt(12))
    para(doc, "{/attachments}")
    para(doc, "{/has_attachments}")


def add_table_loop(doc: Any, key: str, headers: list[str], inner_keys: list[str], widths_mm: list[int] | None = None) -> Any:
    """単純な行繰返し: 繰返し行の先頭セルに {#key}、最終セルに {/key}。縦結合は使わない。"""
    assert len(headers) == len(inner_keys) >= 2
    table = doc.add_table(rows=2, cols=len(headers))
    table.style = doc.styles["Table Grid"]
    table.autofit = False
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        c.paragraphs[0].alignment = CENTER
    last = len(inner_keys) - 1
    for i, k in enumerate(inner_keys):
        c = table.rows[1].cells[i]
        c.text = ""
        txt = "{%s}" % k
        if i == 0:
            txt = "{#%s}" % key + txt
        if i == last:
            txt = txt + "{/%s}" % key
        r = c.paragraphs[0].add_run(txt)
        r.font.size = Pt(10)
    if widths_mm:
        for row in table.rows:
            for i, w in enumerate(widths_mm):
                row.cells[i].width = Mm(w)
    para(doc, "")
    return table


def letter_head(doc: Any, title: str, split_tag: str = "{recipient_name}", with_recipient_address: bool = False) -> None:
    para(doc, "{issue_date_wareki}", align=RIGHT)
    para(doc, "{recipient_organization}")
    if with_recipient_address:
        para(doc, "〒{recipient_postcode}　{recipient_address}")
    if split_tag == "{recipient_name}":
        split_tag_paragraph(doc, "", "{recipient_name}", "{recipient_honorific}", split_at=6)
    else:
        para(doc, "{recipient_name}{recipient_honorific}")
    para(doc, "")
    para(doc, "〒{office_postcode}", align=RIGHT)
    para(doc, "{office_address}", align=RIGHT)
    para(doc, "{office_name}", align=RIGHT)
    para(doc, "電話　{office_phone}　ＦＡＸ　{office_fax}", align=RIGHT)
    para(doc, "弁護士　{lawyer_name}", align=RIGHT)
    para(doc, "")
    para(doc, title, align=CENTER, bold=True, size=16, space_after=Pt(8))
    para(doc, "件名：{subject}")
    if split_tag == "{matter_title}":
        split_tag_paragraph(doc, "案件：", "{matter_title}", split_at=8)
    else:
        para(doc, "案件：{matter_title}")
    para(doc, "（依頼者：{client_name}　／　相手方：{opponent_name}）", size=10)
    para(doc, "")


def court_head(doc: Any, title: str, *, position_tag: str | None = None, show_department: bool = True, split_tag: str = "{case_number}") -> None:
    para(doc, title, align=CENTER, bold=True, size=16, space_after=Pt(8))
    if split_tag == "{case_number}":
        split_tag_paragraph(doc, "事件番号　", "{case_number}", split_at=6)
    else:
        para(doc, "事件番号　{case_number}")
    para(doc, "事件名　　{matter_title}")
    para(doc, "")
    para(doc, "{issue_date_wareki}", align=RIGHT)
    para(doc, "{court_name}" + ("　{court_department}" if show_department else ""))
    para(doc, "{recipient_organization}", size=10)
    para(doc, "{recipient_name}{recipient_honorific}")
    para(doc, "")
    para(doc, "〒{office_postcode}　{office_address}", align=RIGHT)
    para(doc, "{office_name}（電話　{office_phone}　ＦＡＸ　{office_fax}）", align=RIGHT)
    if position_tag:
        para(doc, position_tag + "訴訟代理人弁護士　{lawyer_name}", align=RIGHT)
    else:
        para(doc, "弁護士　{lawyer_name}", align=RIGHT)
    para(doc, "")


def labeled(doc: Any, label: str, tag: str) -> None:
    para(doc, f"{label}：{tag}")


def section_heading(doc: Any, text: str) -> None:
    para(doc, text, bold=True, space_before=Pt(8), space_after=Pt(2))


def closing(doc: Any) -> None:
    para(doc, "")
    para(doc, "以上", align=RIGHT)


def signature_lines(doc: Any, label_tag: str) -> None:
    para(doc, label_tag, space_before=Pt(10))
    para(doc, "住所　{address}", indent_left=Pt(24))
    para(doc, "{representative}", indent_left=Pt(24))
    para(doc, "署名又は記名押印　＿＿＿＿＿＿＿＿＿＿＿＿＿＿　㊞", indent_left=Pt(24))


# ---------------------------------------------------------------------------
# 20 書式の定義
# ---------------------------------------------------------------------------
class TemplateDef:
    def __init__(self, kind: str, title: str, scope: str, recipient_types: list[str], fields: list[dict[str, Any]], *, attachments: bool,
                 build: Callable[[Any], None], values: dict[str, Any], ast: dict[str, Any], sample_attachments: list[dict[str, str]],
                 notes: str, honorific_keys: list[str] | None = None, recipient: dict[str, str] | None = None) -> None:
        self.kind = kind
        self.title = title
        self.scope = scope
        self.recipient_types = recipient_types
        self.fields = fields
        self.attachments = attachments
        self.build = build
        self.values = values
        self.ast = ast
        self.sample_attachments = sample_attachments if attachments else []
        self.notes = notes
        self.honorific_keys = honorific_keys if honorific_keys is not None else ["recipient_name"]
        self.recipient = recipient

    def schema(self) -> dict[str, Any]:
        return {
            "schemaVersion": 1,
            "kind": self.kind,
            "title": self.title,
            "scope": self.scope,
            "recipientTypes": self.recipient_types,
            "fields": self.fields,
            "bodyTag": "body_text",
            "attachments": self.attachments,
            "honorificKeys": self.honorific_keys,
            "notes": self.notes,
        }


def common_letter_fields(*extra: dict[str, Any], recipient_address: bool = False) -> list[dict[str, Any]]:
    fs = [*OFFICE_FIELDS, *RECIPIENT_FIELDS]
    if recipient_address:
        fs += [src("recipient_postcode", "recipient", "recipient.postcode", "宛先郵便番号"), src("recipient_address", "recipient", "recipient.address", "宛先住所")]
    fs += [*MATTER_CORE, SUBJECT, *extra]
    return fs


def common_letter_values(recipient: dict[str, str], subject: str, extra: dict[str, Any], recipient_address: bool = False) -> dict[str, Any]:
    v: dict[str, Any] = {**OFFICE_VALUES, **recipient_values(recipient), **MATTER_VALUES, "subject": subject}
    if recipient_address:
        v["recipient_postcode"] = recipient["postcode"]
        v["recipient_address"] = recipient["address"]
    v.update(extra)
    return v


def common_court_fields(*extra: dict[str, Any], department: bool = True) -> list[dict[str, Any]]:
    fs = [*OFFICE_FIELDS, *RECIPIENT_FIELDS, src("matter_title", "matter", "matter.title", "事件名（案件名）", required=True), CASE_NUMBER, COURT_NAME]
    if department:
        fs.append(COURT_DEPT)
    fs += list(extra)
    return fs


def common_court_values(extra: dict[str, Any], department: bool = True) -> dict[str, Any]:
    v: dict[str, Any] = {**OFFICE_VALUES, **recipient_values(COURT_RCPT), "matter_title": MATTER["title"], "case_number": MATTER["case_number"], "court_name": MATTER["court"]}
    if department:
        v["court_department"] = MATTER["court_department"]
    v.update(extra)
    return v


ATT_LETTER = [{"name": "委任状（写し）", "copies": "１通"}, {"name": "資格証明書（写し）", "copies": "１通"}]
ATT_COURT = [{"name": "訴訟委任状", "copies": "１通"}, {"name": "資格証明書", "copies": "２通"}, {"name": "甲号証写し", "copies": "各２通"}]


def make_templates() -> list[TemplateDef]:
    T: list[TemplateDef] = []

    # ---- T01 受任通知 ----------------------------------------------------
    def b01(doc: Any) -> None:
        letter_head(doc, "受任通知書")
        labeled(doc, "受任範囲", "{engagement_scope}")
        para(doc, "")
        add_body_blocks(doc)
        para(doc, "")
        labeled(doc, "今後の連絡先・連絡方法", "{contact_instruction}")
        loop_block(doc, "acceptance_confirmed", lambda: para(doc, "※受任確認が完了していません。本書面は草案としてのみ出力できます（正式出力不可）。", bold=True), inverted=True)
        add_attachments_block(doc)
        closing(doc)

    T.append(TemplateDef(
        "T01_engagement_notice", "受任通知", "negotiation", ["opponent_agent", "opponent"],
        common_letter_fields(
            src("engagement_scope", "matter", "matter.scope", "受任範囲", "multiline", required=True),
            inp("contact_instruction", "今後の連絡先・連絡方法", "multiline"),
            src("acceptance_confirmed", "matter", "matter.acceptance_confirmed", "受任確認済み", "boolean", required=True, description="案件の受任確認状態（acceptance_status=confirmed）から差し込む。受任確認済みが正式版の条件。未確認のとき書面に草案注記を出す。"),
        ),
        attachments=True, build=b01,
        values=common_letter_values(AGENT_A, "受任のご通知（サンプル）", {
            "engagement_scope": MATTER["scope"],
            "contact_instruction": "本件に関するご連絡は、今後は当職宛にお願いいたします（サンプル）。",
            "acceptance_confirmed": False,
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER)
            .heading("第１　通知の趣旨（サンプル）")
            .numbered("１", "当職は、株式会社ミドリ商事から委任を受け、同社の代理人として本件に関する交渉を担当することになりましたので、その旨ご通知いたします（サンプル文）。")
            .numbered("２", "本件に関するご連絡は、今後は当職宛にお願いいたします（サンプル文）。")
            .heading("第２　お願い（サンプル）")
            .para("本件の事実関係および当方の主張については、別途ご連絡いたします。本サンプルは文例であり、具体的な主張は担当弁護士が記載します。")
            .build(),
        sample_attachments=ATT_LETTER,
        notes="受任確認済み（acceptance_confirmed）が正式版の条件。相手方代理人は所属・住所を併記して選択し、同姓同名を氏名だけで同一人物としない。",
        recipient=AGENT_A,
    ))

    # ---- T02 回答書 --------------------------------------------------------
    def b02(doc: Any) -> None:
        letter_head(doc, "回答書")
        para(doc, "貴職からの「{reference_document_title}」（{reference_document_date_raw}付）について、下記のとおり回答いたします。")
        labeled(doc, "回答方針（弁護士が選択・既定なし）", "{reply_policy}")
        para(doc, "")
        add_table_loop(doc, "reply_items", ["番号", "回答項目", "回答"], ["no", "item", "reply_text"], [15, 50, 95])
        add_body_blocks(doc)
        add_attachments_block(doc)
        closing(doc)

    T.append(TemplateDef(
        "T02_reply_letter", "回答書", "negotiation", ["opponent_agent", "opponent"],
        common_letter_fields(
            inp("reference_document_title", "回答対象の文書名", required=True),
            inp("reference_document_date_raw", "回答対象の文書日付（原文）", "date"),
            legal("reply_policy", "回答方針", "enum", required=True, options=opts("未入力", "回答を留保する", "請求には応じられない", "一部について協議に応じる", "自由記載"), description="既定の認諾等を入れない。弁護士が選択する。"),
            arr("reply_items", "回答項目", [item("no", "番号", required=True, max_graphemes=10), item("item", "回答項目", required=True), item("reply_text", "回答本文", "multiline", required=True)]),
        ),
        attachments=True, build=b02,
        values=common_letter_values(AGENT_B, "ご照会への回答（サンプル）", {
            "reference_document_title": "ご通知書（サンプル）",
            "reference_document_date_raw": date_value("令和８年９月２４日", "2026-09-24"),
            "reply_policy": "未入力",
            "reply_items": [
                {"no": "１", "item": "請求金額について（サンプル）", "reply_text": "（サンプル）ご指摘の金額については、当方で資料を確認中です。確認結果を踏まえ、改めて回答いたします。"},
                {"no": "２", "item": "納入時期について（サンプル）", "reply_text": "（サンプル）納入時期に関する事実関係は、依頼者に確認のうえ別途ご連絡いたします。"},
            ],
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER)
            .para("上記の各回答項目に加え、補足事項を以下に記載します（サンプル）。")
            .numbered("１", "回答方針は本書式では自動選択されません。弁護士が内容を確認して選択します（サンプル文）。")
            .build(),
        sample_attachments=[{"name": "参考資料（写し）", "copies": "１通"}],
        notes="回答方針（reply_policy）に既定の認諾等を入れない。答弁書（T13）とは別種類。",
        recipient=AGENT_B,
    ))

    # ---- T03 照会書 --------------------------------------------------------
    def b03(doc: Any) -> None:
        letter_head(doc, "照会書")
        labeled(doc, "照会の目的", "{inquiry_purpose}")
        para(doc, "")
        add_table_loop(doc, "questions", ["番号", "照会事項"], ["no", "question"], [15, 145])
        labeled(doc, "回答希望日（人が指定）", "{reply_due_date_raw}")
        para(doc, "")
        add_body_blocks(doc)
        closing(doc)

    T.append(TemplateDef(
        "T03_inquiry_letter", "照会書", "negotiation", ["opponent_agent", "opponent", "third_party"],
        common_letter_fields(
            inp("inquiry_purpose", "照会目的", "multiline", required=True),
            arr("questions", "質問", [item("no", "番号", required=True, max_graphemes=10), item("question", "照会事項", "multiline", required=True)]),
            inp("reply_due_date_raw", "回答希望日（原文・人が指定）", "date", required=True),
        ),
        attachments=False, build=b03,
        values=common_letter_values(THIRD_RCPT, "取引経過に関するご照会（サンプル）", {
            "inquiry_purpose": "（サンプル）本件売買取引の経過を確認するため、下記事項について照会いたします。",
            "questions": [
                {"no": "１", "question": "（サンプル）令和７年４月から同年９月までの間に、貴会が受領した資料の有無。"},
                {"no": "２", "question": "（サンプル）上記資料の作成者及び作成日。"},
            ],
            "reply_due_date_raw": date_value("令和８年１０月３１日", "2026-10-31"),
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER).para("ご多忙のところ恐縮ですが、上記回答希望日までにご回答いただけますと幸いです（サンプル）。").build(),
        sample_attachments=[],
        notes="回答希望日は人が指定する（自動算定しない）。第三者宛てでは長い名称の折返しを確認する。",
        recipient=THIRD_RCPT,
    ))

    # ---- T04 資料提出依頼書 -----------------------------------------------
    def b04(doc: Any) -> None:
        letter_head(doc, "資料提出のお願い")
        add_body_blocks(doc)
        para(doc, "")
        add_table_loop(doc, "requested_items", ["番号", "必要資料", "対象期間"], ["no", "item", "period"], [15, 95, 50])
        labeled(doc, "対象期間", "{target_period_raw}")
        labeled(doc, "提出方法", "{submission_method}")
        labeled(doc, "提出希望日", "{due_date_raw}")
        add_attachments_block(doc, "同封書類")
        closing(doc)

    T.append(TemplateDef(
        "T04_document_request", "資料提出依頼書", "client", ["client", "third_party"],
        common_letter_fields(
            arr("requested_items", "必要資料", [item("no", "番号", required=True, max_graphemes=10), item("item", "資料名", required=True), item("period", "対象期間")]),
            inp("target_period_raw", "対象期間（全体）"),
            inp("submission_method", "提出方法"),
            inp("due_date_raw", "提出希望日（原文）", "date"),
        ),
        attachments=True, build=b04,
        values=common_letter_values(CLIENT_RCPT, "資料ご提出のお願い（サンプル）", {
            "requested_items": [
                {"no": "１", "item": "売買基本契約書（写し）", "period": "―"},
                {"no": "２", "item": "請求書・納品書（写し）", "period": "令和７年４月～同年９月"},
                {"no": "３", "item": "入金記録（通帳写し等）", "period": "令和７年４月～令和８年９月"},
            ],
            "target_period_raw": "令和７年４月１日～令和８年９月３０日",
            "submission_method": "郵送またはＰＤＦ（メール添付）",
            "due_date_raw": date_value("令和８年１０月２０日", "2026-10-20"),
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER).para("本件の検討のため、下記の資料のご提出をお願いいたします（サンプル）。").build(),
        sample_attachments=[{"name": "返信用封筒", "copies": "１通"}],
        notes="依頼者または第三者に対する資料提出依頼。希望日は人が指定する。",
        recipient=CLIENT_RCPT,
    ))

    # ---- T05 催告書 --------------------------------------------------------
    def b05(doc: Any) -> None:
        letter_head(doc, "催告書")
        labeled(doc, "請求内容", "{claim_description}")
        labeled(doc, "請求金額（原文）", "{claim_amount}")
        labeled(doc, "根拠の参照", "{claim_basis_ref}")
        labeled(doc, "期限（原文）", "{payment_deadline_raw}")
        labeled(doc, "支払先", "{payment_destination}")
        para(doc, "")
        add_body_blocks(doc)
        closing(doc)

    T.append(TemplateDef(
        "T05_demand_letter", "催告書", "negotiation", ["opponent", "opponent_agent"],
        common_letter_fields(
            inp("claim_description", "請求内容", "multiline", required=True),
            inp("claim_amount", "請求金額（原文/解析値）", "money", required=True),
            inp("claim_basis_ref", "根拠の参照", "multiline"),
            inp("payment_deadline_raw", "期限（原文）", "date", required=True),
            inp("payment_destination", "支払先", "multiline"),
        ),
        attachments=False, build=b05,
        values=common_letter_values(OPPONENT_RCPT, "売買代金のお支払いの催告（サンプル）", {
            "claim_description": "（サンプル）令和７年４月１日付売買基本契約に基づく商品代金のお支払い。",
            "claim_amount": money_value("金３，３００，０００円", "3300000"),
            "claim_basis_ref": "（サンプル）売買基本契約書第３条、令和７年８月３１日付請求書（№０１２）",
            "payment_deadline_raw": date_value("令和８年１０月３１日", "2026-10-31"),
            "payment_destination": "（サンプル）○○銀行○○支店　普通　１２３４５６７　カ）ミドリショウジ",
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER)
            .numbered("１", "当職は、株式会社ミドリ商事の代理人として、上記請求内容のとおりお支払いを求めます（サンプル文）。")
            .numbered("２", "上記期限までにお支払いが確認できない場合の対応については、依頼者と協議のうえ検討いたします（サンプル文。利息・時効等の計算は本書式では自動算定しません）。")
            .build(),
        sample_attachments=[],
        notes="金額は原文（raw）と解析値（decimal）を分けて保持し、原文を差し込む。利息・時効を自動算定しない。期限は原文を保持。",
        recipient=OPPONENT_RCPT,
    ))

    # ---- T06 和解提案書 ----------------------------------------------------
    def b06(doc: Any) -> None:
        letter_head(doc, "和解のご提案")
        add_body_blocks(doc)
        para(doc, "")
        section_heading(doc, "提案条件")
        add_table_loop(doc, "proposal_terms", ["番号", "提案条件"], ["no", "term"], [15, 145])
        labeled(doc, "支払等の条件", "{payment_terms}")
        labeled(doc, "回答希望日", "{reply_due_date_raw}")
        add_attachments_block(doc)
        closing(doc)

    T.append(TemplateDef(
        "T06_settlement_proposal", "和解提案書", "negotiation", ["opponent_agent", "opponent"],
        common_letter_fields(
            arr("proposal_terms", "提案条件", [item("no", "番号", required=True, max_graphemes=10), item("term", "条件", "multiline", required=True)]),
            inp("payment_terms", "支払等の条件", "multiline"),
            inp("reply_due_date_raw", "回答希望日（原文）", "date", required=True),
        ),
        attachments=True, build=b06,
        values=common_letter_values(AGENT_A, "和解のご提案（サンプル）", {
            "proposal_terms": [
                {"no": "１", "term": "（サンプル）解決金の額および支払方法について協議する。"},
                {"no": "２", "term": "（サンプル）本件に関する今後の取引条件について別途協議する。"},
            ],
            "payment_terms": "（サンプル）支払期日・振込先は合意成立後に別途定める。",
            "reply_due_date_raw": date_value("令和８年１１月１４日", "2026-11-14"),
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER).para("本件の早期解決のため、下記のとおり和解をご提案いたします。ご検討のうえ、回答希望日までにご回答いただければ幸いです（サンプル）。").build(),
        sample_attachments=[{"name": "和解案（たたき台）", "copies": "１通"}],
        notes="提案条件は配列。回答希望日は人が指定。",
        recipient=AGENT_A,
    ))

    # ---- T07 合意書 --------------------------------------------------------
    def b07(doc: Any) -> None:
        para(doc, "{agreement_title}", align=CENTER, bold=True, size=16, space_after=Pt(8))
        split_tag_paragraph(doc, "（案件：", "{matter_title}", "）", split_at=8)
        para(doc, "")
        add_body_blocks(doc)
        para(doc, "")
        add_table_loop(doc, "clauses", ["条", "条項"], ["clause_no", "clause_text"], [20, 140])
        labeled(doc, "履行条件", "{performance_terms}")
        para(doc, "")
        para(doc, "本合意の成立を証するため、本書を当事者の数だけ作成し、各自署名又は記名押印のうえ各１通を保有する。")
        para(doc, "{signature_date_raw}", align=RIGHT)
        loop_block(doc, "agreement_parties", lambda: signature_lines(doc, "{label}　{name}"))
        para(doc, "")
        para(doc, "※署名欄の画像や押印画像を電子署名として扱わない（サンプル書式）。", size=9)

    T.append(TemplateDef(
        "T07_agreement", "合意書", "negotiation", ["none", "opponent_agent", "opponent"],
        [
            src("matter_title", "matter", "matter.title", "案件名", required=True),
            inp("agreement_title", "表題", required=True, maxGraphemes=50),
            arr("agreement_parties", "当事者", [item("label", "略称（甲/乙等）", required=True, max_graphemes=10), item("name", "氏名・名称", required=True), item("address", "住所"), item("representative", "法人代表者")], required=True, max_items=20),
            arr("clauses", "条項", [item("clause_no", "条番号", required=True, max_graphemes=20), item("clause_text", "条項本文", "multiline", required=True)], required=True),
            inp("performance_terms", "履行条件", "multiline"),
            inp("signature_date_raw", "署名日（原文）", "date"),
        ],
        attachments=False, build=b07,
        values={
            "matter_title": MATTER["title"],
            "agreement_title": "合意書（サンプル）",
            "agreement_parties": [
                {"label": "甲", "name": MATTER["client_name"], "address": CLIENT_ADDRESS, "representative": CLIENT_REP},
                {"label": "乙", "name": MATTER["opponent_name"], "address": OPPONENT_ADDRESS, "representative": OPPONENT_REP},
            ],
            "clauses": [
                {"clause_no": "第１条", "clause_text": "（サンプル）甲及び乙は、本件取引に関する協議の経過を相互に確認する。"},
                {"clause_no": "第２条", "clause_text": "（サンプル）本合意の各条項の内容は、当事者双方が確認のうえ定める。"},
                {"clause_no": "第３条", "clause_text": "（サンプル）本合意に定めのない事項は、甲乙協議のうえ解決する。"},
            ],
            "performance_terms": "（サンプル）各条項の履行期日・方法は、別紙により定める。",
            "signature_date_raw": date_value("令和８年１０月６日", "2026-10-06"),
        },
        ast=AstBuilder().para(DEMO_DISCLAIMER).para("甲及び乙は、案件表示の件に関し、以下のとおり合意する（サンプル前文）。").build(),
        sample_attachments=[],
        notes="当事者・条項・履行条件・署名欄。署名画像を電子署名と扱わない。宛先は任意（none 可）。",
        honorific_keys=[],
    ))

    # ---- T08 送付書 --------------------------------------------------------
    def b08(doc: Any) -> None:
        letter_head(doc, "送付書", with_recipient_address=True)
        labeled(doc, "送付目的", "{sending_purpose}")
        labeled(doc, "送付方法", "{sending_method}")
        para(doc, "")
        para(doc, "下記の書類を送付いたしますので、ご査収ください。")
        add_attachments_block(doc, "送付書類（書類名　通数）")
        labeled(doc, "連絡事項", "{remarks}")
        add_body_blocks(doc)
        closing(doc)

    T.append(TemplateDef(
        "T08_cover_letter", "送付書", "negotiation", ["opponent_agent", "opponent", "court", "client", "third_party"],
        common_letter_fields(
            inp("sending_purpose", "送付目的", required=True),
            inp("sending_method", "送付方法"),
            inp("remarks", "連絡事項", "multiline"),
            recipient_address=True,
        ),
        attachments=True, build=b08,
        values=common_letter_values(AGENT_A, "書類送付のご案内（サンプル）", {
            "sending_purpose": "（サンプル）受任通知書及び関係資料の送付",
            "sending_method": "郵送（レターパック）",
            "remarks": "（サンプル）ご不明な点がございましたら当職までご連絡ください。",
        }, recipient_address=True),
        ast=AstBuilder().para("（サンプル）以上、よろしくお願い申し上げます。").build(),
        sample_attachments=[{"name": "受任通知書", "copies": "１通"}, {"name": "資料一覧（サンプル）", "copies": "１通"}],
        notes="送付書類名と通数は添付リスト（attachments）で管理する。宛先住所を併記。",
        recipient=AGENT_A,
    ))

    # ---- T09 依頼者への進捗報告書 ----------------------------------------
    def b09(doc: Any) -> None:
        letter_head(doc, "進捗報告書")
        labeled(doc, "報告対象期間", "{report_period_raw}")
        para(doc, "")
        section_heading(doc, "１　経過")
        add_body_blocks(doc)
        section_heading(doc, "２　次の予定")
        para(doc, "{next_steps}", indent_left=Pt(12))
        section_heading(doc, "３　ご依頼事項（ご対応のお願い）")
        para(doc, "{client_requests}", indent_left=Pt(12))
        add_attachments_block(doc, "添付資料")
        closing(doc)

    T.append(TemplateDef(
        "T09_client_progress_report", "依頼者への進捗報告書", "client", ["client"],
        common_letter_fields(
            inp("report_period_raw", "報告対象期間（原文）", required=True),
            inp("next_steps", "次の予定", "multiline"),
            inp("client_requests", "依頼事項", "multiline"),
        ),
        attachments=True, build=b09,
        values=common_letter_values(CLIENT_RCPT, "進捗のご報告（サンプル）", {
            "report_period_raw": "令和８年９月１日～同年９月３０日",
            "next_steps": "（サンプル）相手方代理人からの回答を待ち、回答内容を踏まえて今後の方針を協議する予定です。",
            "client_requests": "（サンプル）入金記録（通帳写し）のご提出をお願いいたします。",
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER)
            .numbered("（１）", "令和８年９月１０日、相手方代理人宛に受任通知書を送付しました（サンプル）。")
            .numbered("（２）", "同月２４日、相手方代理人から回答書を受領しました。内容は現在検討中です（サンプル）。")
            .build(),
        sample_attachments=[{"name": "相手方回答書（写し）", "copies": "１通"}],
        notes="報告対象期間・経過（本文）・次の予定・依頼事項。",
        recipient=CLIENT_RCPT,
    ))

    # ---- T10 委任契約書 ----------------------------------------------------
    def b10(doc: Any) -> None:
        para(doc, "委任契約書", align=CENTER, bold=True, size=16, space_after=Pt(8))
        split_tag_paragraph(doc, "委任者（甲）　", "{client_name}", split_at=7)
        para(doc, "住所　{recipient_address}", indent_left=Pt(24))
        para(doc, "受任者（乙）　{office_name}　弁護士　{lawyer_name}")
        para(doc, "〒{office_postcode}　{office_address}（電話　{office_phone}　ＦＡＸ　{office_fax}）", indent_left=Pt(24))
        para(doc, "")
        para(doc, "甲及び乙は、案件「{matter_title}」に関し、次のとおり委任契約を締結する。")
        add_body_blocks(doc)
        section_heading(doc, "第１条（委任の範囲）")
        para(doc, "{delegation_scope}", indent_left=Pt(12))
        section_heading(doc, "第２条（報酬）")
        para(doc, "{fee_terms}", indent_left=Pt(12))
        section_heading(doc, "第３条（費用）")
        para(doc, "{expense_terms}", indent_left=Pt(12))
        section_heading(doc, "第４条（契約の終了）")
        para(doc, "{termination_terms}", indent_left=Pt(12))
        para(doc, "")
        para(doc, "{contract_date_raw}", align=RIGHT)
        para(doc, "甲　{client_name}　　署名又は記名押印　＿＿＿＿＿＿＿＿＿＿　㊞", space_before=Pt(10))
        para(doc, "乙　{office_name}　弁護士　{lawyer_name}　　＿＿＿＿＿＿＿＿＿＿　㊞", space_before=Pt(6))
        para(doc, "")
        para(doc, "※報酬率や法的適合性を本書式が自動保証するものではない（サンプル書式）。", size=9)

    T.append(TemplateDef(
        "T10_retainer_agreement", "委任契約書", "client", ["client"],
        [
            *OFFICE_FIELDS,
            src("recipient_address", "recipient", "recipient.address", "委任者住所"),
            src("matter_title", "matter", "matter.title", "案件名", required=True),
            src("client_name", "matter", "matter.client_name", "委任者（依頼者）", required=True),
            inp("delegation_scope", "委任範囲", "multiline", required=True),
            inp("fee_terms", "報酬条件", "multiline", required=True, description="報酬率や法的適合性を自動保証しない。所内書式の文言を弁護士が確認する。"),
            inp("expense_terms", "費用", "multiline"),
            inp("termination_terms", "終了条件等", "multiline"),
            inp("contract_date_raw", "契約日（原文）", "date", required=True),
        ],
        attachments=False, build=b10,
        values={
            **OFFICE_VALUES,
            "recipient_address": CLIENT_ADDRESS,
            "matter_title": MATTER["title"],
            "client_name": MATTER["client_name"],
            "delegation_scope": "（サンプル）髙﨑産業株式会社に対する売買代金請求に関する交渉及び第一審訴訟手続。",
            "fee_terms": "（サンプル）着手金・報酬金の額及び支払時期は、所内報酬基準に基づき別紙のとおりとする。",
            "expense_terms": "（サンプル）実費（印紙・郵券・交通費等）は甲の負担とし、乙の請求に基づき精算する。",
            "termination_terms": "（サンプル）甲又は乙は、相手方に通知して本契約を終了させることができる。終了時の精算方法は別紙による。",
            "contract_date_raw": date_value("令和８年１０月６日", "2026-10-06"),
        },
        ast=AstBuilder().para(DEMO_DISCLAIMER).build(),
        sample_attachments=[],
        notes="当事者・委任範囲・報酬条件・費用・終了条件の所内書式。報酬率や法的適合性を自動保証しない。",
        honorific_keys=[],
        recipient=CLIENT_RCPT,
    ))

    # ---- T11 訴訟委任状 ----------------------------------------------------
    AUTH_ITEMS = [
        ("auth_counterclaim", "反訴の提起"),
        ("auth_withdrawal", "訴えの取下げ"),
        ("auth_settlement", "和解"),
        ("auth_waiver_admission", "請求の放棄・認諾"),
        ("auth_appeal", "控訴、上告、上告受理の申立て及びこれらの取下げ"),
        ("auth_sub_delegation", "復代理人の選任"),
    ]

    def b11(doc: Any) -> None:
        para(doc, "訴訟委任状", align=CENTER, bold=True, size=16, space_after=Pt(8))
        para(doc, "{poa_date_raw}", align=RIGHT)
        para(doc, "{recipient_organization}", size=10)
        para(doc, "{recipient_name}{recipient_honorific}")
        para(doc, "")
        split_tag_paragraph(doc, "委任者　", "{principal_name}", split_at=6)
        para(doc, "住所　{principal_address}", indent_left=Pt(24))
        para(doc, "")
        para(doc, "私は、次の弁護士を訴訟代理人と定め、下記事件に関する各事項を委任します。")
        para(doc, "受任弁護士　{office_name}　弁護士　{lawyer_name}", indent_left=Pt(12))
        para(doc, "〒{office_postcode}　{office_address}（電話　{office_phone}　ＦＡＸ　{office_fax}）", indent_left=Pt(24))
        section_heading(doc, "事件の表示")
        para(doc, "裁判所　　{court_name}", indent_left=Pt(12))
        para(doc, "事件番号　{case_number}", indent_left=Pt(12))
        para(doc, "事件名　　{matter_title}", indent_left=Pt(12))
        para(doc, "相手方　　{opponent_name}", indent_left=Pt(12))
        section_heading(doc, "委任事項")
        para(doc, "１　上記事件に関する一切の訴訟行為", indent_left=Pt(12))
        para(doc, "２　上記事件に関する書類の受領及び提出", indent_left=Pt(12))
        section_heading(doc, "特別授権事項（選択したものに限る。既定では選択なし・弁護士確認必須）")
        for key, label in AUTH_ITEMS:
            loop_block(doc, key, lambda label=label: para(doc, "一　" + label, indent_left=Pt(12)))
        para(doc, "{special_auth_other}", indent_left=Pt(12))
        add_body_blocks(doc)
        para(doc, "")
        para(doc, "委任者　{principal_name}　　署名又は記名押印　＿＿＿＿＿＿＿＿＿＿　㊞", space_before=Pt(10))

    T.append(TemplateDef(
        "T11_power_of_attorney_litigation", "訴訟委任状", "court", ["court", "none"],
        [
            *OFFICE_FIELDS,
            *RECIPIENT_FIELDS,
            src("principal_name", "matter", "matter.client_name", "委任者", required=True),
            inp("principal_address", "委任者住所", required=True),
            src("matter_title", "matter", "matter.title", "事件名（案件名）", required=True),
            CASE_NUMBER,
            COURT_NAME,
            src("opponent_name", "matter", "matter.opponent_name", "相手方", required=True),
            *[legal(k, f"特別授権：{label}", "boolean", description="既定 OFF。弁護士確認必須。自動で ON にしない。") for k, label in AUTH_ITEMS],
            legal("special_auth_other", "特別授権：その他（自由記載）", "multiline"),
            inp("poa_date_raw", "委任日（原文）", "date", required=True),
        ],
        attachments=False, build=b11,
        values={
            **OFFICE_VALUES,
            **recipient_values(COURT_RCPT),
            "principal_name": MATTER["client_name"],
            "principal_address": CLIENT_ADDRESS,
            "matter_title": MATTER["title"],
            "case_number": MATTER["case_number"],
            "court_name": MATTER["court"],
            "opponent_name": MATTER["opponent_name"],
            **{k: False for k, _ in AUTH_ITEMS},
            "special_auth_other": "",
            "poa_date_raw": date_value("令和８年１０月６日", "2026-10-06"),
        },
        ast=AstBuilder().para("（サンプル）特別授権事項は既定で選択されていません。委任者の意思を確認した弁護士が選択します。").build(),
        sample_attachments=[],
        notes="特別授権の選択は既定 OFF（legalJudgement）。弁護士確認必須。委任者・受任弁護士・対象事件・権限条項。",
        recipient=COURT_RCPT,
    ))

    # ---- T12 訴状 ----------------------------------------------------------
    def b12(doc: Any) -> None:
        para(doc, "訴　　状", align=CENTER, bold=True, size=18, space_after=Pt(8))
        para(doc, "{issue_date_wareki}", align=RIGHT)
        para(doc, "{court_name}　{recipient_organization}", size=10)
        para(doc, "{recipient_name}{recipient_honorific}")
        para(doc, "")
        para(doc, "〒{office_postcode}　{office_address}", align=RIGHT)
        para(doc, "{office_name}（電話　{office_phone}　ＦＡＸ　{office_fax}）", align=RIGHT)
        para(doc, "原告訴訟代理人弁護士　{lawyer_name}", align=RIGHT)
        para(doc, "")
        section_heading(doc, "当事者の表示")
        para(doc, "〒　{plaintiff_address}", indent_left=Pt(12))
        split_tag_paragraph(doc, "原告　", "{plaintiff_name}", split_at=7)
        para(doc, "〒　{defendant_address}", indent_left=Pt(12))
        para(doc, "被告　{defendant_name}")
        para(doc, "")
        para(doc, "{case_name}", align=CENTER, bold=True)
        para(doc, "（案件：{matter_title}）", align=CENTER, size=9)
        para(doc, "訴訟物の価額　{claim_value_raw}")
        para(doc, "貼用印紙額　　{stamp_fee_raw}")
        para(doc, "予納郵券　　　{postage_raw}")
        section_heading(doc, "第１　請求の趣旨")
        loop_block(doc, "claim_purpose_items", lambda: para(doc, "{no}　{text}", indent_left=Pt(12)))
        section_heading(doc, "第２　請求の原因")
        add_body_blocks(doc)
        section_heading(doc, "証拠方法")
        loop_block(doc, "evidence_refs", lambda: para(doc, "{label}　{title}", indent_left=Pt(12)))
        add_attachments_block(doc, "附属書類")

    T.append(TemplateDef(
        "T12_complaint", "訴状", "court", ["court"],
        [
            *OFFICE_FIELDS, *RECIPIENT_FIELDS,
            src("matter_title", "matter", "matter.title", "案件名", required=True),
            COURT_NAME,
            PLAINTIFF, inp("plaintiff_address", "原告住所", required=True),
            DEFENDANT, inp("defendant_address", "被告住所", required=True),
            inp("case_name", "事件名", required=True, maxGraphemes=60),
            inp("claim_value_raw", "訴額（人が入力・確認）", required=True, description="自動計算しない。原文のまま保持。"),
            inp("stamp_fee_raw", "貼用印紙額（人が入力・確認）", required=True, description="自動計算しない。原文のまま保持。"),
            inp("postage_raw", "予納郵券（原文）"),
            legal("claim_purpose_items", "請求の趣旨", "array", itemFields=[item("no", "番号", required=True, max_graphemes=10), item("text", "請求の趣旨", "multiline", required=True)], maxItems=50, required=True),
            EVIDENCE_REFS,
        ],
        attachments=True, build=b12,
        values={k: v for k, v in common_court_values({
            "plaintiff_name": MATTER["client_name"], "plaintiff_address": CLIENT_ADDRESS + "　" + CLIENT_REP,
            "defendant_name": MATTER["opponent_name"], "defendant_address": OPPONENT_ADDRESS + "　" + OPPONENT_REP,
            "case_name": "売買代金請求事件",
            "claim_value_raw": "金３，３００，０００円（人が入力・確認）",
            "stamp_fee_raw": "金２２，０００円（人が入力・確認）",
            "postage_raw": "（人が入力）",
            "claim_purpose_items": [
                {"no": "１", "text": "（サンプル）被告は、原告に対し、金３，３００，０００円及びこれに対する遅延損害金を支払え。"},
                {"no": "２", "text": "（サンプル）訴訟費用は被告の負担とする。"},
                {"no": "３", "text": "（サンプル）との判決並びに仮執行の宣言を求める。"},
            ],
            "evidence_refs": [{"label": "甲第１号証", "title": "売買基本契約書（サンプル）"}, {"label": "甲第２号証", "title": "請求書（サンプル）"}, {"label": "甲第３号証の１", "title": "納品書（サンプル・枝番）"}],
        }, department=False).items() if k != "case_number"},
        ast=AstBuilder().para(DEMO_DISCLAIMER)
            .heading("１　当事者（サンプル）")
            .numbered("（１）", "原告は、食品卸売業等を目的とする株式会社である（サンプル）。")
            .numbered("（２）", "被告は、食品小売業等を目的とする株式会社である（サンプル）。")
            .heading("２　売買契約の締結（サンプル）")
            .para("原告は、令和７年４月１日、被告との間で売買基本契約を締結したと主張する（甲１）。契約の具体的内容は、担当弁護士が証拠に基づき記載する（サンプル）。")
            .heading("３　まとめ（サンプル）")
            .para("よって、原告は被告に対し、請求の趣旨記載の判決を求める旨を記載する（サンプル・結論部分は弁護士が記載）。")
            .build(),
        sample_attachments=ATT_COURT,
        notes="訴額・印紙額は人が入力・確認（自動計算しない）。請求の趣旨は legalJudgement。証拠方法は evidence_refs。原告/被告の表示は人が指定する。",
        recipient=COURT_RCPT,
    ))

    # ---- T13 答弁書 --------------------------------------------------------
    def b13(doc: Any) -> None:
        court_head(doc, "答弁書", position_tag="被告")
        para(doc, "原告　{plaintiff_name}")
        para(doc, "被告　{defendant_name}")
        section_heading(doc, "第１　請求の趣旨に対する答弁")
        loop_block(doc, "answer_items", lambda: para(doc, "{no}　{text}", indent_left=Pt(12)))
        section_heading(doc, "第２　請求の原因に対する認否（段落単位・自動選択しない）")
        add_table_loop(doc, "admissions", ["段落", "請求原因の記載", "認否", "自由記載"], ["para_no", "claim_text", "admission", "admission_note"], [15, 75, 25, 45])
        section_heading(doc, "第３　被告の主張")
        add_body_blocks(doc)
        section_heading(doc, "証拠方法")
        loop_block(doc, "evidence_refs", lambda: para(doc, "{label}　{title}", indent_left=Pt(12)))
        add_attachments_block(doc, "附属書類")

    T.append(TemplateDef(
        "T13_answer", "答弁書", "court", ["court"],
        common_court_fields(
            PLAINTIFF, DEFENDANT,
            legal("answer_items", "請求の趣旨に対する答弁", "array", itemFields=[item("no", "番号", required=True, max_graphemes=10), item("text", "答弁", "multiline", required=True)], maxItems=50, required=True),
            legal("admissions", "請求原因に対する認否（段落単位）", "array", itemFields=[
                item("para_no", "段落番号", required=True, max_graphemes=20),
                item("claim_text", "請求原因の記載（原文）", "multiline", required=True),
                item("admission", "認否", "enum", required=True, options=["未入力", "認める", "否認", "不知", "自由記載"]),
                item("admission_note", "自由記載・理由", "multiline"),
            ], maxItems=500, required=True, description="段落単位で未入力/認める/否認/不知/自由記載を選べるが自動選択しない。"),
            EVIDENCE_REFS,
        ),
        attachments=True, build=b13,
        values=common_court_values({
            "plaintiff_name": MATTER["opponent_name"], "defendant_name": MATTER["client_name"],
            "answer_items": [{"no": "１", "text": "（サンプル・弁護士が記載）"}, {"no": "２", "text": "（サンプル・弁護士が記載）"}],
            "admissions": [
                {"para_no": "１（１）", "claim_text": "（サンプル）原告は食品卸売業等を目的とする株式会社である。", "admission": "未入力", "admission_note": ""},
                {"para_no": "１（２）", "claim_text": "（サンプル）被告は食品小売業等を目的とする株式会社である。", "admission": "未入力", "admission_note": ""},
                {"para_no": "２", "claim_text": "（サンプル）原告と被告は令和７年４月１日に売買基本契約を締結した。", "admission": "未入力", "admission_note": ""},
            ],
            "evidence_refs": [{"label": "乙第１号証", "title": "（サンプル）取引明細"}],
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER).para("被告の主張は、追って準備書面において詳述する予定である（サンプル・内容未定）。").build(),
        sample_attachments=[{"name": "訴訟委任状", "copies": "１通"}, {"name": "乙号証写し", "copies": "各２通"}],
        notes="認否は段落単位の enum（未入力/認める/否認/不知/自由記載）で legalJudgement。サンプルでは未入力のまま。回答書（T02）とは別種類。",
        recipient=COURT_RCPT,
    ))

    # ---- T14 準備書面 ------------------------------------------------------
    def b14(doc: Any) -> None:
        para(doc, "準備書面（{brief_number_raw}）", align=CENTER, bold=True, size=16, space_after=Pt(8))
        split_tag_paragraph(doc, "事件番号　", "{case_number}", split_at=6)
        para(doc, "事件名　　{matter_title}")
        para(doc, "原告　{plaintiff_name}　／　被告　{defendant_name}")
        para(doc, "")
        para(doc, "{issue_date_wareki}", align=RIGHT)
        para(doc, "{court_name}　{court_department}")
        para(doc, "{recipient_organization}", size=10)
        para(doc, "{recipient_name}{recipient_honorific}")
        para(doc, "")
        para(doc, "〒{office_postcode}　{office_address}", align=RIGHT)
        para(doc, "{office_name}（電話　{office_phone}　ＦＡＸ　{office_fax}）", align=RIGHT)
        para(doc, "{party_position}訴訟代理人弁護士　{lawyer_name}", align=RIGHT)
        para(doc, "")
        add_body_blocks(doc)
        section_heading(doc, "証拠参照")
        loop_block(doc, "evidence_refs", lambda: para(doc, "{label}　{title}", indent_left=Pt(12)))
        add_attachments_block(doc, "附属書類")

    T.append(TemplateDef(
        "T14_brief", "準備書面", "court", ["court"],
        common_court_fields(
            PLAINTIFF, DEFENDANT,
            inp("brief_number_raw", "書面番号（原文）", required=True, maxGraphemes=20, description="「第１」「１」等を文字列として保持。案件内の衝突検査はアプリ側で行う。"),
            inp("party_position", "提出者の立場", "enum", required=True, options=opts("原告", "被告", "申立人", "相手方", "その他")),
            EVIDENCE_REFS,
        ),
        attachments=True, build=b14,
        values=common_court_values({
            "plaintiff_name": MATTER["client_name"], "defendant_name": MATTER["opponent_name"],
            "brief_number_raw": "第１", "party_position": "原告",
            "evidence_refs": [{"label": "甲第３号証の１", "title": "納品書（サンプル・枝番）"}, {"label": "甲第３号証の２", "title": "納品書（サンプル・枝番）"}],
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER)
            .heading("第１　本書面の目的（サンプル）")
            .para("本書面は、被告答弁書に対する原告の反論の骨子を整理するものである（サンプル）。")
            .heading("第２　争点に関する主張（サンプル）")
            .numbered("１", "納品の事実について、原告は納品書（甲３の１・甲３の２）を提出して主張する予定である（サンプル）。")
            .numbered("（１）", "納品日・数量等の具体的内容は、証拠に基づき担当弁護士が記載する（サンプル）。", level=2)
            .build(),
        sample_attachments=[{"name": "甲第３号証の１・２（写し）", "copies": "各２通"}],
        notes="書面番号は案件内で衝突検査（アプリ側）。主張の見出し/段落は本文 AST。証拠参照は evidence_refs。",
        recipient=COURT_RCPT,
    ))

    # ---- T15 当事者目録 ----------------------------------------------------
    def b15(doc: Any) -> None:
        para(doc, "当事者目録", align=CENTER, bold=True, size=16, space_after=Pt(8))
        para(doc, "事件番号　{case_number}")
        para(doc, "事件名　　{matter_title}")
        split_tag_paragraph(doc, "裁判所　　", "{court_name}", split_at=7)
        para(doc, "{issue_date_wareki}", align=RIGHT)
        para(doc, "{recipient_organization}", size=10)
        para(doc, "{recipient_name}{recipient_honorific}")
        para(doc, "")
        add_table_loop(doc, "parties", ["立場", "氏名・名称", "住所", "法人代表者", "代理人"], ["role_label", "name", "address", "representative", "agent"], [18, 36, 46, 30, 30])
        para(doc, "〒{office_postcode}　{office_address}", align=RIGHT)
        para(doc, "作成者　{office_name}　弁護士　{lawyer_name}（電話　{office_phone}　ＦＡＸ　{office_fax}）", align=RIGHT)
        add_body_blocks(doc)

    T.append(TemplateDef(
        "T15_party_list", "当事者目録", "court", ["court"],
        [
            *OFFICE_FIELDS, *RECIPIENT_FIELDS,
            src("matter_title", "matter", "matter.title", "事件名（案件名）", required=True),
            CASE_NUMBER, COURT_NAME,
            arr("parties", "当事者", [
                item("role_label", "立場（原告/被告等）", required=True, max_graphemes=20),
                item("name", "氏名・名称", required=True),
                item("address", "住所", "multiline"),
                item("representative", "法人代表者"),
                item("agent", "代理人（複数は改行区切り）", "multiline"),
            ], required=True, max_items=50, description="複数当事者対応。代理人は当事者ごとに記載し、氏名一致で自動補完しない。"),
        ],
        attachments=False, build=b15,
        values=common_court_values({
            "parties": [
                {"role_label": "原告", "name": MATTER["client_name"], "address": CLIENT_ADDRESS, "representative": CLIENT_REP, "agent": "原告訴訟代理人弁護士　青葉　花子（青葉法律事務所）"},
                {"role_label": "被告", "name": MATTER["opponent_name"], "address": OPPONENT_ADDRESS, "representative": OPPONENT_REP, "agent": "被告訴訟代理人弁護士　佐藤　一郎（さくら総合法律事務所）"},
                {"role_label": "被告", "name": LONG_ORG_NAME, "address": THIRD_RCPT["address"], "representative": "代表理事　渡邊　﨑太郎", "agent": "被告訴訟代理人弁護士　佐藤　一郎（港中央法律事務所）"},
            ],
        }, department=False),
        ast=AstBuilder().para("（サンプル）当事者が多数の場合は行を追加する。長い名称は縮小せず、折返しを確認する。").build(),
        sample_attachments=[],
        notes="当事者配列・立場・住所・法人代表者・代理人。単純な行繰返し（縦結合なし）。同姓同名の代理人は所属で区別。",
        recipient=COURT_RCPT,
    ))

    # ---- T16 証拠説明書 ----------------------------------------------------
    def b16(doc: Any) -> None:
        court_head(doc, "証拠説明書", position_tag="{submitter_position}")
        add_table_loop(doc, "evidence_rows", ["号証", "標目", "作成年月日", "作成者", "原本・写し", "立証趣旨"], ["label", "title", "made_date", "author", "original_copy", "proof_purpose"], [22, 36, 24, 24, 16, 38])
        para(doc, "※立証趣旨は弁護士が入力する。証拠番号の原文（全角・枝番）は変更しない。", size=9)
        add_body_blocks(doc)

    T.append(TemplateDef(
        "T16_evidence_description", "証拠説明書", "court", ["court"],
        common_court_fields(
            SUBMITTER_POSITION,
            arr("evidence_rows", "証拠", [
                item("label", "証拠番号（表示ラベル）", required=True, max_graphemes=50),
                item("title", "標目", required=True),
                item("made_date", "作成日（原文）"),
                item("author", "作成者"),
                item("original_copy", "原本/写し", "enum", required=True, options=["原本", "写し"]),
                item("proof_purpose", "立証趣旨（弁護士が入力）", "multiline", required=True),
            ], required=True, max_items=500),
        ),
        attachments=False, build=b16,
        values=common_court_values({
            "submitter_position": "原告",
            "evidence_rows": [
                {"label": "甲第１号証", "title": "売買基本契約書（サンプル）", "made_date": "令和７年４月１日", "author": "原告・被告", "original_copy": "写し", "proof_purpose": "（弁護士が入力・サンプル）原告と被告の間の取引条件"},
                {"label": "甲第２号証", "title": "請求書（サンプル）", "made_date": "令和７年８月３１日", "author": "原告", "original_copy": "写し", "proof_purpose": "（弁護士が入力・サンプル）請求の事実"},
                {"label": "甲第３号証の１", "title": "納品書（サンプル・枝番）", "made_date": "令和７年５月１０日", "author": "原告", "original_copy": "原本", "proof_purpose": "（弁護士が入力・サンプル）納品の事実"},
            ],
        }),
        ast=AstBuilder().para("（サンプル）備考欄がある場合はここに記載する。").build(),
        sample_attachments=[],
        notes="証拠番号・標目・作成日・作成者・原本/写し・立証趣旨。立証趣旨は弁護士が入力。表は単純な行繰返し。",
        recipient=COURT_RCPT,
    ))

    # ---- T17 電子証拠説明書 ------------------------------------------------
    def b17(doc: Any) -> None:
        court_head(doc, "電子証拠説明書（ファイル対応表）", position_tag="{submitter_position}")
        labeled(doc, "書式プロファイル", "{format_profile}")
        labeled(doc, "提出媒体・方法", "{media_note}")
        para(doc, "")
        add_table_loop(doc, "electronic_evidence_rows", ["号証", "標目", "ファイル名", "形式", "頁数", "ハッシュ等", "立証趣旨"], ["label", "title", "file_name", "file_format", "page_count", "hash_note", "proof_purpose"], [20, 28, 32, 14, 12, 28, 26])
        para(doc, "※本書式は形式上の確認のみを目的とする汎用プロファイル。特定の裁判所・手続の電子提出仕様への完全対応を表示しない。", size=9)
        add_body_blocks(doc)

    T.append(TemplateDef(
        "T17_electronic_evidence_description", "電子証拠説明書", "court", ["court"],
        common_court_fields(
            SUBMITTER_POSITION,
            inp("format_profile", "書式プロファイル", required=True, description="選択した書式プロファイル名。T16 と同一スキーマにしない。"),
            inp("media_note", "提出媒体・方法"),
            arr("electronic_evidence_rows", "電子証拠（ファイル対応）", [
                item("label", "証拠番号（表示ラベル）", required=True, max_graphemes=50),
                item("title", "標目", required=True),
                item("file_name", "ファイル名（表示用原文）", required=True),
                item("file_format", "ファイル形式", "enum", required=True, options=["PDF", "PDF（OCRあり）", "画像", "その他"]),
                item("page_count", "頁数（原文）"),
                item("hash_note", "ハッシュ・版の注記"),
                item("proof_purpose", "立証趣旨（弁護士が入力）", "multiline", required=True),
            ], required=True, max_items=500),
        ),
        attachments=False, build=b17,
        values=common_court_values({
            "submitter_position": "原告",
            "format_profile": "generic_electronic_v1（形式確認のみ）",
            "media_note": "（サンプル）ＰＤＦファイルを所定の方法で提出",
            "electronic_evidence_rows": [
                {"label": "甲第１号証", "title": "売買基本契約書（サンプル）", "file_name": "甲０１＿売買基本契約書.pdf", "file_format": "PDF", "page_count": "１２", "hash_note": "sha256: 0000…demo（サンプル）", "proof_purpose": "（弁護士が入力・サンプル）取引条件"},
                {"label": "甲第３号証の１", "title": "納品書（サンプル・枝番）", "file_name": "甲０３－１＿納品書.pdf", "file_format": "PDF（OCRあり）", "page_count": "１", "hash_note": "sha256: 1111…demo（サンプル）", "proof_purpose": "（弁護士が入力・サンプル）納品の事実"},
            ],
        }),
        ast=AstBuilder().para("（サンプル）ファイル名は表示用の原文を保持し、保存キーは別に管理する。").build(),
        sample_attachments=[],
        notes="電子提出用のファイル対応情報（ファイル名・形式・頁数・ハッシュ）を持つ別スキーマ。T16 の表題だけ変えた同一スキーマにしない。mints 等への完全対応を表示しない。",
        recipient=COURT_RCPT,
    ))

    # ---- T18 書証提出書 ----------------------------------------------------
    def b18(doc: Any) -> None:
        court_head(doc, "書証提出書", position_tag="{submitter_position}")
        para(doc, "下記の書証を提出します。")
        add_table_loop(doc, "submitted_evidence", ["号証", "標目", "通数"], ["label", "title", "copies"], [30, 100, 30])
        loop_block(doc, "applicability_confirmed", lambda: para(doc, "※選択された手続・裁判所にこの書式を適用できるか未確認です（草案のみ）。", bold=True), inverted=True)
        add_body_blocks(doc)
        add_attachments_block(doc, "対応資料")

    T.append(TemplateDef(
        "T18_documentary_evidence_submission", "書証提出書", "court", ["court"],
        common_court_fields(
            SUBMITTER_POSITION,
            arr("submitted_evidence", "提出する書証", [item("label", "証拠番号（表示ラベル）", required=True, max_graphemes=50), item("title", "標目", required=True), item("copies", "通数（原文）", required=True, max_graphemes=20)], required=True, max_items=500),
            inp("applicability_confirmed", "手続・裁判所への適用を確認済み", "boolean", required=True, description="選択された手続/裁判所に適用するか確認する。既定 OFF。"),
        ),
        attachments=True, build=b18,
        values=common_court_values({
            "submitter_position": "原告",
            "submitted_evidence": [{"label": "甲第１号証", "title": "売買基本契約書（サンプル）", "copies": "２通"}, {"label": "甲第２号証", "title": "請求書（サンプル）", "copies": "２通"}, {"label": "甲第３号証の１", "title": "納品書（サンプル・枝番）", "copies": "２通"}],
            "applicability_confirmed": False,
        }),
        ast=AstBuilder().para("（サンプル）提出に関する連絡事項があればここに記載する。").build(),
        sample_attachments=[{"name": "証拠説明書", "copies": "１通"}],
        notes="事件情報・提出する書証一覧・対応資料。選択された手続/裁判所に適用するか確認（applicability_confirmed）。",
        recipient=COURT_RCPT,
    ))

    # ---- T19 訴え取下書 ----------------------------------------------------
    def b19(doc: Any) -> None:
        court_head(doc, "訴え取下書", position_tag="原告")
        para(doc, "原告　{plaintiff_name}")
        para(doc, "被告　{defendant_name}")
        para(doc, "")
        labeled(doc, "取下げの範囲", "{withdrawal_scope}")
        para(doc, "{withdrawal_scope_detail}", indent_left=Pt(12))
        labeled(doc, "対象当事者", "{target_parties}")
        labeled(doc, "相手方の同意（弁護士確認）", "{consent_status}")
        para(doc, "")
        add_body_blocks(doc)
        para(doc, "※取下げに必要な同意等の有無・要否の判断は本書式が自動化しない（サンプル書式）。", size=9)

    T.append(TemplateDef(
        "T19_withdrawal_of_action", "訴え取下書", "court", ["court"],
        common_court_fields(
            PLAINTIFF, DEFENDANT,
            legal("withdrawal_scope", "取下げの範囲", "enum", required=True, options=opts("未入力", "訴えの全部", "訴えの一部", "自由記載")),
            legal("withdrawal_scope_detail", "取下げ範囲の詳細（一部・自由記載）", "multiline", visibilityRule={"op": "in", "field": "withdrawal_scope", "values": ["訴えの一部", "自由記載"]}, requiredRule={"op": "in", "field": "withdrawal_scope", "values": ["訴えの一部", "自由記載"]}),
            inp("target_parties", "対象当事者", "multiline", required=True),
            legal("consent_status", "相手方の同意（弁護士確認）", "enum", required=True, options=opts("未確認", "同意を取得済み（同意書添付）", "同意不要と判断（弁護士確認）", "自由記載"), description="必要な同意等の判断を自動化しない。"),
        ),
        attachments=False, build=b19,
        values=common_court_values({
            "plaintiff_name": MATTER["client_name"], "defendant_name": MATTER["opponent_name"],
            "withdrawal_scope": "未入力", "withdrawal_scope_detail": "",
            "target_parties": "（サンプル）被告　髙﨑産業株式会社",
            "consent_status": "未確認",
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER).para("原告は、頭書事件について、上記範囲の訴えを取り下げる旨を記載する（サンプル・取下げ範囲と同意の要否は弁護士が確認）。").build(),
        sample_attachments=[],
        notes="事件情報・取下げの範囲・対象当事者・本文。必要な同意等の判断を自動化しない（consent_status は legalJudgement）。",
        recipient=COURT_RCPT,
    ))

    # ---- T20 委任終了通知 --------------------------------------------------
    def b20(doc: Any) -> None:
        letter_head(doc, "委任終了のご通知")
        labeled(doc, "対象事件", "{matter_title}")
        labeled(doc, "終了日", "{termination_date_raw}")
        para(doc, "")
        add_body_blocks(doc)
        section_heading(doc, "引継事項")
        loop_block(doc, "handover_items", lambda: para(doc, "{no}　{item}", indent_left=Pt(12)))
        labeled(doc, "書類の返還等", "{document_return_note}")
        add_attachments_block(doc, "返還・同封書類")
        closing(doc)

    T.append(TemplateDef(
        "T20_termination_of_engagement", "委任終了通知", "client", ["client", "opponent_agent", "court"],
        common_letter_fields(
            inp("termination_date_raw", "終了日（原文）", "date", required=True),
            arr("handover_items", "引継事項", [item("no", "番号", required=True, max_graphemes=10), item("item", "引継事項", "multiline", required=True)]),
            inp("document_return_note", "書類の返還等", "multiline"),
        ),
        attachments=True, build=b20,
        values=common_letter_values(CLIENT_RCPT, "委任終了のご通知（サンプル）", {
            "termination_date_raw": date_value("令和８年１０月３１日", "2026-10-31"),
            "handover_items": [{"no": "１", "item": "（サンプル）相手方代理人との連絡経過の一覧"}, {"no": "２", "item": "（サンプル）お預かりした原本資料の返還"}],
            "document_return_note": "（サンプル）お預かりしている書類は、受領書と引換えに返還いたします。",
        }),
        ast=AstBuilder().para(DEMO_DISCLAIMER).para("上記対象事件に関する委任は、上記終了日をもって終了いたします（サンプル。終了の理由・経緯は担当弁護士が記載）。案件の記録はアプリ上で自動削除されません。").build(),
        sample_attachments=[{"name": "受領書（返送用）", "copies": "１通"}],
        notes="対象事件・終了日・通知先・本文・引継事項。案件を自動削除しない。",
        recipient=CLIENT_RCPT,
    ))

    assert len(T) == 20, len(T)
    return T


# ---------------------------------------------------------------------------
# DOCX 後処理・検査
# ---------------------------------------------------------------------------
def rezip_deterministic(data: bytes) -> bytes:
    """entry 順を固定し、タイムスタンプを 1980-01-01 に揃えて再圧縮する（冪等性のため）。"""
    with zipfile.ZipFile(io.BytesIO(data)) as zin:
        entries = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    order = sorted(entries, key=lambda n: (n != "[Content_Types].xml", n))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
        for n in order:
            zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zout.writestr(zi, entries[n])
    return out.getvalue()


def docx_bytes(doc: Any) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return rezip_deterministic(buf.getvalue())


class TagScan:
    def __init__(self) -> None:
        self.tags: list[tuple[str, str, str, str]] = []  # (kind, key, raw, location)
        self.errors: list[str] = []
        self.split_run_tags: list[str] = []


def scan_part_tags(xml: bytes, part_name: str, scan: TagScan) -> None:
    """段落単位で w:t を連結し、{...} を抽出する（tags.ts の scanParagraphTags 相当）。段落跨ぎ・不均衡を検出。"""
    root = etree.fromstring(xml)
    stack: list[tuple[str, str]] = []
    for pi, p in enumerate(root.iter(W + "p")):
        texts = [t.text or "" for t in p.iter(W + "t")]
        text = "".join(texts)
        loc = f"{part_name}/p{pi + 1}"
        # run 境界オフセット（分割 run の検出用）
        boundaries = set()
        acc = 0
        for t in texts[:-1]:
            acc += len(t)
            boundaries.add(acc)
        i = 0
        while i < len(text):
            o = text.find("{", i)
            c = text.find("}", i)
            if o < 0:
                if c >= 0:
                    scan.errors.append(f"{loc}: stray '}}'")
                break
            if 0 <= c < o:
                scan.errors.append(f"{loc}: stray '}}' before '{{'")
            c = text.find("}", o + 1)
            if c < 0:
                scan.errors.append(f"{loc}: unclosed '{{' (paragraph-spanning tag?) in {text!r}")
                break
            inner = text[o + 1:c]
            raw = "{" + inner + "}"
            kind = "value"
            key = inner
            if inner.startswith("#"):
                kind, key = "loop_open", inner[1:]
            elif inner.startswith("/"):
                kind, key = "loop_close", inner[1:]
            elif inner.startswith("^"):
                kind, key = "inverted_open", inner[1:]
            if not TAG_KEY_RE.match(key) or key in DANGEROUS_KEYS:
                scan.errors.append(f"{loc}: invalid tag {raw}")
            if any(o < b < c + 1 for b in boundaries):
                scan.split_run_tags.append(raw)
            if kind in ("loop_open", "inverted_open"):
                stack.append((key, loc))
            elif kind == "loop_close":
                if not stack or stack[-1][0] != key:
                    scan.errors.append(f"{loc}: unbalanced loop close {raw}")
                else:
                    stack.pop()
            scan.tags.append((kind, key, raw, loc))
            i = c + 1
    for key, loc in stack:
        scan.errors.append(f"{loc}: loop {{#{key}}} not closed")


def inspect_docx(data: bytes, tdef: TemplateDef) -> dict[str, Any]:
    """生成した DOCX を検査: 開ける / 構成 part / タグ集合とスキーマの整合 / 分割 run タグ / 単一セクション。"""
    report: dict[str, Any] = {"errors": [], "warnings": []}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        report["entries"] = len(names)
        doc_parts = [n for n in names if n == "word/document.xml"]
        headers = [n for n in names if re.match(r"word/header\d*\.xml$", n)]
        footers = [n for n in names if re.match(r"word/footer\d*\.xml$", n)]
        if len(doc_parts) != 1:
            report["errors"].append(f"document.xml count={len(doc_parts)}")
        if len(headers) < 1 or len(footers) < 1:
            report["errors"].append(f"header/footer missing: {headers} {footers}")
        for bad in ("word/comments.xml", "word/footnotes.xml", "word/endnotes.xml", "word/vbaProject.bin"):
            if bad in names:
                report["errors"].append(f"forbidden part present: {bad}")
        if any(n.startswith("customXml/") for n in names) or "docProps/thumbnail.jpeg" in names:
            report["warnings"].append("optional parts not stripped")
        ct = z.read("[Content_Types].xml").decode("utf-8")
        if "macroEnabled" in ct:
            report["errors"].append("macroEnabled content type")
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
        if 'TargetMode="External"' in rels:
            report["errors"].append("external relationship in document.xml.rels")
        docxml = z.read("word/document.xml")
        if docxml.count(b"<w:sectPr") != 1:
            report["errors"].append(f"sectPr count={docxml.count(b'<w:sectPr')}")
        for needle in (b"w:txbxContent", b"<w:ins ", b"<w:del ", b"w:object", b"footnoteReference", b"commentRangeStart"):
            if needle in docxml:
                report["errors"].append(f"forbidden element {needle!r}")
        for m in re.finditer(rb'w:instr="([^"]*)"', docxml + b"".join(z.read(n) for n in headers + footers)):
            if m.group(1).strip() != b"PAGE":
                report["errors"].append(f"unsupported field {m.group(1)!r}")
        scan = TagScan()
        scan_part_tags(docxml, "document", scan)
        for h in headers:
            scan_part_tags(z.read(h), h, scan)
        for f in footers:
            scan_part_tags(z.read(f), f, scan)
        report["errors"] += scan.errors
        report["tags"] = len(scan.tags)
        report["split_run_tags"] = scan.split_run_tags
        if not scan.split_run_tags:
            report["errors"].append("no split-run tag found")
        header_text = b"".join(z.read(h) for h in headers)
        footer_text = b"".join(z.read(f) for f in footers)
        if b"{document_status_label}" not in header_text:
            report["errors"].append("header lacks {document_status_label}")
        if b"{template_version_label}" not in footer_text or b'w:instr="PAGE"' not in footer_text:
            report["errors"].append("footer lacks PAGE field or {template_version_label}")
        # スタイル: Noto Serif JP の rFonts 明示
        styles = z.read("word/styles.xml").decode("utf-8")
        if f'w:eastAsia="{FONT}"' not in styles or f'w:ascii="{FONT}"' not in styles or f'w:hAnsi="{FONT}"' not in styles:
            report["errors"].append("rFonts ascii/hAnsi/eastAsia not set to Noto Serif JP")
        # スキーマ整合
        fields = {f["key"]: f for f in tdef.fields}
        item_keys = {it["key"] for f in tdef.fields if f["type"] == "array" for it in f["itemFields"]}
        used = {key for _, key, _, _ in scan.tags}
        for kind, key, raw, loc in scan.tags:
            if key in fields or key in SYSTEM_SUPPLIED_KEYS or key in item_keys or key == "body_text":
                continue
            report["errors"].append(f"{loc}: undefined tag {raw}")
        for key, f in fields.items():
            if key not in used:
                report["errors"].append(f"field {key} has no tag in DOCX")
        if "body_blocks" not in used:
            report["errors"].append("no {#body_blocks}")
        if tdef.attachments and "has_attachments" not in used:
            report["errors"].append("attachments=true but no {#has_attachments}")
        if not tdef.attachments and "has_attachments" in used:
            report["errors"].append("attachments=false but {#has_attachments} used")
        loops = sorted({key for kind, key, _, _ in scan.tags if kind in ("loop_open", "inverted_open")})
        report["loops"] = loops
    # python-docx で再度開けること
    d = Document(io.BytesIO(data))
    report["paragraphs"] = len(d.paragraphs)
    report["tables"] = len(d.tables)
    sec = d.sections[0]
    if abs(sec.page_width.mm - 210) > 0.5 or abs(sec.page_height.mm - 297) > 0.5 or len(d.sections) != 1:
        report["errors"].append("page is not single-section A4 portrait")
    return report


def validate_schema_shape(schema: dict[str, Any]) -> list[str]:
    """TemplateFieldSchemaSchema の主要制約を Python 側でも確認する（zod による本検証は node 側）。"""
    errs: list[str] = []
    keys: set[str] = set()
    for f in schema["fields"]:
        k = f["key"]
        if k in keys:
            errs.append(f"duplicate key {k}")
        keys.add(k)
        if not TAG_KEY_RE.match(k) or k in DANGEROUS_KEYS or "." in k:
            errs.append(f"bad key {k}")
        if f["source"] != "input" and not f.get("sourcePath"):
            errs.append(f"{k}: source without sourcePath")
        if f["type"] == "enum" and not f.get("options"):
            errs.append(f"{k}: enum without options")
        if f["type"] == "array" and not f.get("itemFields"):
            errs.append(f"{k}: array without itemFields")
        if f["type"] != "array" and f.get("itemFields"):
            errs.append(f"{k}: itemFields on non-array")
        if f.get("legalJudgement") and f.get("defaultPolicy", "none") != "none":
            errs.append(f"{k}: legalJudgement with default")
        if f.get("legalJudgement") and f.get("defaultValue") is not None:
            errs.append(f"{k}: legalJudgement with defaultValue")
        for it in f.get("itemFields", []) or []:
            if not TAG_KEY_RE.match(it["key"]) or it["key"] in DANGEROUS_KEYS:
                errs.append(f"{k}.{it['key']}: bad item key")
            if it["type"] == "enum" and not it.get("options"):
                errs.append(f"{k}.{it['key']}: enum item without options")
        for rule_name in ("visibilityRule", "requiredRule"):
            rule = f.get(rule_name)
            if rule:
                for ref in _condition_fields(rule):
                    if ref not in {x["key"] for x in schema["fields"]}:
                        errs.append(f"{k}: {rule_name} references unknown field {ref}")
    if len(schema["fields"]) < 1 or len(schema["fields"]) > 200:
        errs.append("fields count out of range")
    for hk in schema["honorificKeys"]:
        if hk not in keys:
            errs.append(f"honorificKeys references undefined field {hk}")
    if len(schema.get("notes", "")) > 2000:
        errs.append("notes too long")
    return errs


def _condition_fields(c: dict[str, Any]) -> set[str]:
    op = c["op"]
    if op in ("and", "or"):
        out: set[str] = set()
        for x in c["conditions"]:
            out |= _condition_fields(x)
        return out
    if op == "not":
        return _condition_fields(c["condition"])
    return {c["field"]}


def validate_sample_values(tdef: TemplateDef) -> list[str]:
    errs: list[str] = []
    fields = {f["key"]: f for f in tdef.fields}
    for k in fields:
        if k not in tdef.values:
            errs.append(f"sample value missing for {k}")
    for k, v in tdef.values.items():
        f = fields.get(k)
        if f is None:
            errs.append(f"sample value for unknown field {k}")
            continue
        if f["type"] == "array":
            if not isinstance(v, list):
                errs.append(f"{k}: array sample must be list")
                continue
            item_keys = {it["key"] for it in f["itemFields"]}
            for idx, row in enumerate(v):
                extra = set(row) - item_keys
                if extra:
                    errs.append(f"{k}[{idx}]: unknown item keys {sorted(extra)}")
                for it in f["itemFields"]:
                    if it.get("required") and not row.get(it["key"]):
                        errs.append(f"{k}[{idx}].{it['key']}: required item empty")
                    if it["type"] == "enum" and row.get(it["key"]) not in {o["value"] for o in it["options"]}:
                        errs.append(f"{k}[{idx}].{it['key']}: enum value {row.get(it['key'])!r} not in options")
        elif f["type"] == "boolean":
            if not isinstance(v, bool):
                errs.append(f"{k}: boolean sample must be bool")
            if f.get("legalJudgement") and v is not False:
                errs.append(f"{k}: legalJudgement boolean sample must be OFF")
        elif f["type"] == "enum":
            if v not in {o["value"] for o in f["options"]}:
                errs.append(f"{k}: enum sample {v!r} not in options")
        elif f["type"] == "money":
            if not (isinstance(v, dict) and "raw" in v and re.fullmatch(r"\d+", v.get("decimal", ""))):
                errs.append(f"{k}: money sample needs raw + decimal")
        elif f["type"] == "date":
            if not (isinstance(v, dict) and "raw" in v and "iso" in v and "precision" in v):
                errs.append(f"{k}: date sample needs raw/iso/precision")
        elif f.get("required") and not v:
            errs.append(f"{k}: required sample is empty")
    if tdef.kind == "T13_answer":
        for row in tdef.values["admissions"]:
            if row["admission"] != "未入力":
                errs.append("T13 認否 sample must stay 未入力")
    blocks = tdef.ast["blocks"]
    ids = [b["id"] for b in blocks]
    if len(ids) != len(set(ids)):
        errs.append("AST block ids not unique")
    for b in blocks:
        if not re.fullmatch(r"[A-Za-z0-9_-]+", b["id"]):
            errs.append(f"AST id invalid {b['id']}")
        if "監修済" in b.get("text", ""):
            errs.append("AST text claims 監修済み")
    for text in json.dumps(tdef.values, ensure_ascii=False), tdef.notes:
        if "監修済み" in text and "監修済みではない" not in text:
            errs.append("sample text claims 監修済み")
    return errs


# ---------------------------------------------------------------------------
# 出力
# ---------------------------------------------------------------------------
def dump_json(obj: Any) -> bytes:
    return (json.dumps(obj, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def write_if_changed(path: Path, data: bytes, check_only: bool) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and path.read_bytes() == data:
        return "unchanged"
    if check_only:
        return "DIFFERS"
    path.write_bytes(data)
    return "written"


def provenance(tdef: TemplateDef) -> dict[str, Any]:
    return {
        "source": "RayLaw demo (架空・自作)",
        "sampleOnly": True,
        "supervised": False,
        "license": "demo-internal",
        "sourceUrl": None,
        "checkedAt": CHECKED_AT,
        "applicability": "デモ専用。所内弁護士の確認と承認前は正式出力不可",
        "notes": "監修済みではない",
        "kind": tdef.kind,
        "title": tdef.title,
        "generator": "scripts/build_demo_templates.py",
    }


def sample_values_doc(tdef: TemplateDef) -> dict[str, Any]:
    return {
        "schemaVersion": 1,
        "kind": tdef.kind,
        "sampleOnly": True,
        "note": SAMPLE_NOTE,
        "recipient": tdef.recipient,
        "values": rv(tdef.values),
        "attachments": tdef.sample_attachments,
        "renderContext": {
            "documentStatusLabel": "サンプル・要確認",
            "templateVersionLabel": f"{tdef.kind} demo-r1（sample_only・監修済みではない）",
            "issueDateIso": ISSUE_DATE_ISO,
            "issueDateWareki": ISSUE_DATE_WAREKI,
        },
        "astFile": "ast_sample.json",
    }


def readme_text(templates: list[TemplateDef], reports: dict[str, dict[str, Any]]) -> str:
    lines = [
        "# templates/demo — 架空デモ書式 20 種（T01..T20）",
        "",
        "`scripts/build_demo_templates.py`（python-docx）が生成する。各ディレクトリ名は `TemplateKindSchema` の値。",
        "",
        "初期状態の法務書式は sample_only。草案出力には「サンプル・要確認」を表示。監修済みではない。",
        "",
        "各ディレクトリの構成: `template.docx`（A4 縦・単一セクション・横書き・Noto Serif JP 12pt、ヘッダー `{document_status_label}`、"
        "フッター PAGE フィールド + `{template_version_label}`）、`schema.json`（TemplateFieldSchema）、`provenance.json`"
        "（sampleOnly:true / supervised:false）、`sample_values.json`（架空値）、`ast_sample.json`（本文 DocumentAST）。",
        "",
        "本文は `{#body_blocks}` 段落ループ（`{#is_heading}` / `{#is_numbered}` / `{#is_paragraph}` の各条件も段落単位）で見出し・番号付き段落・通常段落を別段落として出力する。"
        "添付リストは `{#has_attachments}` → `{#attachments}{name}　{copies}{/attachments}` の段落ループ。表は単純な行繰返しのみ（縦結合なし）。"
        "各書式に、意図的に複数 run に分割したタグ（例: `{recip` + `ient_name}`）を 1 つ以上含む。",
        "",
        "| kind | 表題 | scope | recipientTypes | 添付 | loops（array / 条件） | 分割 run タグ |",
        "|---|---|---|---|---|---|---|",
    ]
    for t in templates:
        rep = reports[t.kind]
        loops = ", ".join(f"`{k}`" for k in rep["loops"])
        split = ", ".join(f"`{s}`" for s in rep["split_run_tags"])
        lines.append(f"| `{t.kind}` | {t.title} | {t.scope} | {', '.join(t.recipient_types)} | {'あり' if t.attachments else 'なし'} | {loops} | {split} |")
    lines += [
        "",
        "## 法的判断を含む項目（legalJudgement:true・既定値なし）",
        "",
        "- T02 `reply_policy`（回答方針）: 既定の認諾等を入れない。",
        "- T11 `auth_*`（特別授権）: 既定 OFF。弁護士確認必須。",
        "- T12 `claim_purpose_items`（請求の趣旨）、T13 `answer_items` / `admissions`（認否: 未入力/認める/否認/不知/自由記載。サンプルは未入力のまま）。",
        "- T19 `withdrawal_scope` / `withdrawal_scope_detail` / `consent_status`。",
        "",
        "## 注意",
        "",
        "- 架空の「青葉法律事務所」「青葉　花子」「株式会社ミドリ商事」「髙﨑産業株式会社」および同姓同名の相手方代理人「佐藤　一郎」（所属が異なる 2 名）を使用。実在の人物・団体・事件とは無関係。",
        "- 全角数字・全角空白・異体字（髙・﨑）・30 grapheme 超の名称・証拠枝番「甲第３号証の１」を含む。原文を変換しない。",
        "- 文例は明らかにサンプルであり、法的結論を事実として記載しない。T05 の利息・時効、T12 の訴額・印紙額は自動算定しない。",
        "- 出典・適用範囲・確認日・利用条件は `provenance.json`。裁判所の公開書式は使用していない。",
        "- 再生成: `python3 -I scripts/build_demo_templates.py`（冪等。`--check` で差分検査のみ）。",
        "",
    ]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT, help="出力ディレクトリ（既定 templates/demo）")
    ap.add_argument("--check", action="store_true", help="書き込まず差分の有無だけ報告する")
    args = ap.parse_args(argv)
    out: Path = args.out

    templates = make_templates()
    kinds = [t.kind for t in templates]
    assert len(set(kinds)) == 20

    total_errors = 0
    reports: dict[str, dict[str, Any]] = {}
    statuses: list[tuple[str, dict[str, str]]] = []
    for t in templates:
        doc = new_document(t.title, t.kind)
        t.build(doc)
        data = docx_bytes(doc)
        rep = inspect_docx(data, t)
        schema = t.schema()
        rep["errors"] += validate_schema_shape(schema)
        rep["errors"] += validate_sample_values(t)
        reports[t.kind] = rep
        rep["sha256"] = hashlib.sha256(data).hexdigest()
        d = out / t.kind
        st = {
            "template.docx": write_if_changed(d / "template.docx", data, args.check),
            "schema.json": write_if_changed(d / "schema.json", dump_json(schema), args.check),
            "provenance.json": write_if_changed(d / "provenance.json", dump_json(provenance(t)), args.check),
            "sample_values.json": write_if_changed(d / "sample_values.json", dump_json(sample_values_doc(t)), args.check),
            "ast_sample.json": write_if_changed(d / "ast_sample.json", dump_json(t.ast), args.check),
        }
        statuses.append((t.kind, st))
        total_errors += len(rep["errors"])

    readme_status = write_if_changed(out / "README.md", readme_text(templates, reports).encode("utf-8"), args.check)

    print("== RayLaw demo templates ==")
    print(f"out: {out}")
    print(f"{'kind':40} {'docx':9} {'json':9} entries paras tbls tags loops  split-run tags")
    for kind, st in statuses:
        rep = reports[kind]
        json_status = "unchanged" if all(v == "unchanged" for k, v in st.items() if k != "template.docx") else ("DIFFERS" if any(v == "DIFFERS" for v in st.values()) else "written")
        print(f"{kind:40} {st['template.docx']:9} {json_status:9} {rep['entries']:7d} {rep['paragraphs']:5d} {rep['tables']:4d} {rep['tags']:4d} {len(rep['loops']):5d}  {', '.join(rep['split_run_tags'])}")
        for e in rep["errors"]:
            print(f"    ERROR: {e}")
        for w in rep["warnings"]:
            print(f"    warning: {w}")
    print(f"README.md: {readme_status}")
    n_fields = sum(len(t.fields) for t in templates)
    n_legal = sum(1 for t in templates for f in t.fields if f.get("legalJudgement"))
    n_arrays = sum(1 for t in templates for f in t.fields if f["type"] == "array")
    print(f"templates: {len(templates)}  fields: {n_fields}  array fields: {n_arrays}  legalJudgement fields: {n_legal}  errors: {total_errors}")
    changed = [k for k, st in statuses if any(v != "unchanged" for v in st.values())]
    if readme_status != "unchanged":
        changed.append("README.md")
    print(f"changed: {len(changed)}" + (f" ({', '.join(changed)})" if changed else " (idempotent: nothing to write)"))
    if total_errors:
        return 1
    if args.check and changed:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
